#!/usr/bin/env python3
"""마크다운 보고서 초안 → docx(편집용) + pdf(제출용).

  python3 build_report.py --md draft/전문가활용_260403.md --out out/ --format docx pdf

지원 문법(보고서에 필요한 것만): #/##/### 제목, 불릿(-, *, ●, ○ / 들여쓰기 2단),
번호 목록, **굵게**, 표(|), 인용(>), 구분선(---), `키: 값` 머리글.
PDF 는 Chrome headless 인쇄로 만든다(LibreOffice 불필요).
"""
from __future__ import annotations

import argparse
import html
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    shutil.which("google-chrome") or "",
    shutil.which("chromium") or "",
]
FONT_STACK = '"Apple SD Gothic Neo","맑은 고딕","Malgun Gothic","Noto Sans KR",sans-serif'


# ------------------------------------------------------------------ parsing
def parse(md: str) -> list[dict]:
    blocks: list[dict] = []
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            blocks.append({"t": "hr"})
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            blocks.append({"t": "h", "level": len(m.group(1)), "text": m.group(2)})
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[\s:\-]+$", "".join(cells)):
                    rows.append(cells)
                i += 1
            blocks.append({"t": "table", "rows": rows})
            continue
        if stripped.startswith(">"):
            blocks.append({"t": "quote", "text": stripped.lstrip("> ").strip()})
            i += 1
            continue
        m = re.match(r"^(\s*)(?:[-*•]|●|○|▪|◦)\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            level = 1 if indent < 2 else (2 if indent < 6 else 3)
            if stripped.startswith("○") or stripped.startswith("◦"):
                level = max(level, 2)
            blocks.append({"t": "li", "level": level, "text": m.group(2)})
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            blocks.append({"t": "ol", "level": 1 if indent < 2 else 2,
                           "num": m.group(2), "text": m.group(3)})
            i += 1
            continue
        m = re.match(r"^([가-힣A-Za-z][가-힣A-Za-z0-9 /()]{0,14}):\s+(.+)$", stripped)
        if m and len(blocks) < 8:
            blocks.append({"t": "kv", "key": m.group(1), "val": m.group(2)})
            i += 1
            continue
        blocks.append({"t": "p", "text": stripped})
        i += 1
    return blocks


def inline_runs(text: str):
    """**굵게** 만 처리해 (텍스트, bold) 쌍으로 쪼갠다."""
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True))
        else:
            out.append((part.replace("`", ""), False))
    return out


# ------------------------------------------------------------------ docx
def to_docx(blocks: list[dict], out_path: Path, title_hint: str):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.0)
        section.left_margin = section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(4)

    def para(text_runs, size=11, bold=False, align=None, indent=0.0, space_before=0):
        p = doc.add_paragraph()
        if align:
            p.alignment = align
        pf = p.paragraph_format
        pf.left_indent = Cm(indent)
        pf.space_before = Pt(space_before)
        for text, is_bold in text_runs:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold or is_bold
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        return p

    BULLET = {1: "●", 2: "○", 3: "▪"}
    for b in blocks:
        t = b["t"]
        if t == "h":
            lvl = b["level"]
            if lvl == 1:
                para(inline_runs(b["text"]), size=16, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6)
            else:
                para(inline_runs(b["text"]), size=13 if lvl == 2 else 11.5,
                     bold=True, space_before=10)
        elif t == "kv":
            para([(f"{b['key']}: ", True), (b["val"], False)])
        elif t == "li":
            para([(f"{BULLET.get(b['level'], '·')} ", False)] + inline_runs(b["text"]),
                 indent=0.5 * b["level"])
        elif t == "ol":
            para([(f"{b['num']}. ", False)] + inline_runs(b["text"]), indent=0.5 * b["level"])
        elif t == "quote":
            p = para(inline_runs(b["text"]), indent=0.8)
            p.runs[0].font.italic = True
        elif t == "table":
            rows = b["rows"]
            tbl = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            tbl.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    c = tbl.cell(ri, ci)
                    c.text = re.sub(r"\*\*", "", cell)
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9.5)
                            r.font.name = "맑은 고딕"
                            r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
                            r.font.bold = ri == 0
        elif t == "hr":
            para([("", False)])
        else:
            para(inline_runs(b["text"]))

    doc.save(str(out_path))
    return out_path


# ------------------------------------------------------------------ html/pdf
def to_html(blocks: list[dict], title: str) -> str:
    def esc(s):
        return html.escape(s)

    def inline(text):
        return "".join(
            f"<b>{esc(t)}</b>" if bold else esc(t) for t, bold in inline_runs(text)
        )

    body, list_open = [], None

    def close_list():
        nonlocal list_open
        if list_open:
            body.append(f"</{list_open}>")
            list_open = None

    for b in blocks:
        t = b["t"]
        if t in ("li", "ol"):
            want = "ul" if t == "li" else "ol"
            if list_open != want:
                close_list()
                body.append(f"<{want}>")
                list_open = want
            body.append(f'<li class="l{b["level"]}">{inline(b["text"])}</li>')
            continue
        close_list()
        if t == "h":
            body.append(f'<h{b["level"]}>{inline(b["text"])}</h{b["level"]}>')
        elif t == "kv":
            body.append(f'<p class="kv"><b>{esc(b["key"])}:</b> {esc(b["val"])}</p>')
        elif t == "quote":
            body.append(f"<blockquote>{inline(b['text'])}</blockquote>")
        elif t == "hr":
            body.append("<hr>")
        elif t == "table":
            rows = b["rows"]
            cells = []
            for ri, row in enumerate(rows):
                tag = "th" if ri == 0 else "td"
                cells.append("<tr>" + "".join(f"<{tag}>{inline(c)}</{tag}>" for c in row) + "</tr>")
            body.append("<table>" + "".join(cells) + "</table>")
        else:
            body.append(f"<p>{inline(b['text'])}</p>")
    close_list()

    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8">
<title>{esc(title)}</title><style>
@page {{ size: A4; margin: 20mm 20mm; }}
body {{ font-family: {FONT_STACK}; font-size: 11pt; line-height: 1.75; color: #111; }}
h1 {{ font-size: 16pt; text-align: center; margin: 0 0 18px; letter-spacing: -0.01em; }}
h2 {{ font-size: 12.5pt; margin: 20px 0 6px; }}
h3 {{ font-size: 11.5pt; margin: 14px 0 4px; }}
p {{ margin: 0 0 8px; text-align: justify; }}
p.kv {{ margin: 0 0 3px; }}
ul, ol {{ margin: 0 0 10px; padding-left: 1.2em; }}
li {{ margin: 2px 0; }}
li.l2 {{ margin-left: 1.1em; list-style-type: circle; }}
li.l3 {{ margin-left: 2.2em; list-style-type: square; }}
blockquote {{ margin: 8px 0 8px 1em; padding-left: 10px; border-left: 2px solid #ccc; color: #444; }}
table {{ border-collapse: collapse; width: 100%; margin: 8px 0 14px; font-size: 9.5pt; }}
th, td {{ border: 1px solid #999; padding: 5px 7px; text-align: left; vertical-align: top; }}
th {{ background: #f2f4f7; font-weight: 700; }}
hr {{ border: 0; border-top: 1px solid #ddd; margin: 14px 0; }}
</style></head><body>
{chr(10).join(body)}
</body></html>"""


def to_pdf(html_text: str, out_path: Path) -> Path | None:
    chrome = next((c for c in CHROME_CANDIDATES if c and Path(c).exists()), None)
    if not chrome:
        print("[pdf] Chrome/Chromium 을 찾지 못해 PDF 를 건너뜁니다. docx 를 사용하세요.",
              file=sys.stderr)
        return None
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "doc.html"
        src.write_text(html_text, encoding="utf-8")
        cmd = [chrome, "--headless=new", "--disable-gpu", "--no-sandbox",
               "--no-pdf-header-footer", f"--print-to-pdf={out_path}", src.as_uri()]
        r = subprocess.run(cmd, capture_output=True, text=True)
        if not out_path.exists():
            cmd = [c for c in cmd if c != "--no-pdf-header-footer"]
            cmd[1] = "--headless"
            r = subprocess.run(cmd, capture_output=True, text=True)
        if not out_path.exists():
            print(f"[pdf] 실패: {r.stderr[-400:]}", file=sys.stderr)
            return None
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--format", nargs="+", default=["docx", "pdf"], choices=["docx", "pdf", "html"])
    args = ap.parse_args()

    md_path = Path(args.md)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    blocks = parse(md_path.read_text(encoding="utf-8"))
    title = next((b["text"] for b in blocks if b["t"] == "h" and b["level"] == 1), md_path.stem)
    stem = md_path.stem

    made = []
    if "docx" in args.format:
        made.append(to_docx(blocks, out_dir / f"{stem}.docx", title))
    html_text = to_html(blocks, title)
    if "html" in args.format:
        p = out_dir / f"{stem}.html"
        p.write_text(html_text, encoding="utf-8")
        made.append(p)
    if "pdf" in args.format:
        p = to_pdf(html_text, out_dir / f"{stem}.pdf")
        if p:
            made.append(p)

    for p in made:
        print(f"[ok] {p}")
    if not made:
        sys.exit(1)


if __name__ == "__main__":
    main()
