#!/usr/bin/env python3
"""회의자료 → 텍스트 추출. 실패해도 파이프라인을 막지 않는다(원본은 첨부로 남는다)."""
from __future__ import annotations

import argparse
import html
import re
import subprocess
import sys
from pathlib import Path


def from_pdf(p: Path) -> str:
    try:
        return subprocess.run(["pdftotext", "-layout", str(p), "-"],
                              capture_output=True, text=True, check=True).stdout
    except Exception as e:
        return f"[추출 실패: pdftotext {e}]"


def from_docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    parts = [para.text for para in doc.paragraphs if para.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def from_pptx(p: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(p))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n--- slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"[note] {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(out)


def from_xlsx(p: Path) -> str:
    import warnings
    warnings.filterwarnings("ignore")
    import openpyxl
    wb = openpyxl.load_workbook(str(p), data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n--- sheet: {ws.title} ---")
        for row in ws.iter_rows():
            vals = [str(c.value).strip() for c in row if c.value is not None and str(c.value).strip()]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)


def from_drawio(p: Path) -> str:
    s = p.read_text(encoding="utf-8", errors="replace")
    pages = re.findall(r'<diagram[^>]*name="([^"]*)"', s)
    labels = [html.unescape(v) for v in re.findall(r'value="([^"]*)"', s) if v.strip()]
    labels = [re.sub(r"<[^>]+>", " ", v).strip() for v in labels]
    out = ["[pages] " + ", ".join(pages)]
    out += [f"- {v}" for v in labels if v]
    return "\n".join(out)


def from_plain(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="replace")


HANDLERS = {
    ".pdf": from_pdf, ".docx": from_docx, ".pptx": from_pptx, ".xlsx": from_xlsx,
    ".drawio": from_drawio, ".txt": from_plain, ".md": from_plain, ".csv": from_plain,
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    src = Path(args.input)
    ext = src.suffix.lower()

    if ext in (".hwp", ".hwpx"):
        Path(args.out).write_text(
            f"[본문 미추출] {src.name}\n"
            "hwp/hwpx 텍스트 추출기가 없습니다. 원본은 첨부로 유지됩니다.\n"
            "본문이 필요하면 한글에서 PDF 로 저장한 뒤 다시 인테이크하세요.\n",
            encoding="utf-8")
        print(f"[hwp] 본문 미추출 — 원본만 첨부: {src.name}", file=sys.stderr)
        return

    fn = HANDLERS.get(ext)
    if not fn:
        print(f"[skip] 미지원 형식: {src.name}", file=sys.stderr)
        return
    try:
        text = fn(src)
    except Exception as e:
        text = f"[추출 실패] {src.name}: {e}"
        print(text, file=sys.stderr)
    Path(args.out).write_text(text, encoding="utf-8")
    print(f"[ok] {src.name} → {Path(args.out).name} ({len(text)} chars)")


if __name__ == "__main__":
    main()
